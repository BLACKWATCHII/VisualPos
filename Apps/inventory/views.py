from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.conf import settings
from django.template.loader import render_to_string
from django.http import FileResponse, HttpResponse
from django.urls import reverse
from django.utils import timezone
from .models import InventoryAdjustment, InventoryHistory
from item.models import Item

from datetime import datetime, time
from io import BytesIO
import os
import subprocess
import uuid

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

@login_required
def inventory_adjustment_create(request):
    items = Item.objects.all()

    download_id = request.GET.get('download')
    download_url = None
    if download_id:
        try:
            download_url = reverse('inventory_adjustment_receipt', args=[int(download_id)])
        except (TypeError, ValueError):
            download_url = None

    # Si es POST, registramos el ajuste
    if request.method == 'POST':
        item_id = request.POST.get('item')
        adjustment_type = request.POST.get('adjustment_type')
        quantity = int(request.POST.get('quantity'))
        reason = request.POST.get('reason')

        item = get_object_or_404(Item, id=item_id)
        stock_before = item.Stock

        if adjustment_type == 'OUT' and quantity > item.Stock:
            messages.error(request, "No hay suficiente stock para esta salida.")
            return redirect('inventory_adjustment_create')

        adjustment = InventoryAdjustment.objects.create(
            item=item,
            adjustment_type=adjustment_type,
            quantity=quantity,
            reason=reason,
            user=request.user
        )

        if adjustment_type == 'IN':
            item.Stock += quantity
        else:
            item.Stock -= quantity
        item.save()

        InventoryHistory.objects.create(
            adjustment=adjustment,
            stock_before=stock_before,
            stock_after=item.Stock
        )

        messages.success(request, "Ajuste de inventario registrado correctamente. Descargando comprobante...")
        return redirect(f"{reverse('inventory_adjustment_create')}?download={adjustment.id}")

    # Historial paginado
    adjustments_list = InventoryAdjustment.objects.select_related('item', 'user').order_by('-created_at')
    paginator = Paginator(adjustments_list, 10) 
    page_number = request.GET.get('page')
    adjustments = paginator.get_page(page_number)

    return render(request, 'inventory/inventory_adjustment_form.html', {
        'items': items,
        'adjustments': adjustments,
        'download_url': download_url,
    })


def _render_pdf_inline_with_puppeteer(template_src: str, context: dict, filename: str = "ReciboAjusteInventario.pdf", as_attachment: bool = True):
    html_id = str(uuid.uuid4())
    tmp_dir = os.path.join(settings.BASE_DIR, "tmp")
    html_path = os.path.join(tmp_dir, f"{html_id}.html")
    pdf_path = os.path.join(tmp_dir, f"{html_id}.pdf")

    os.makedirs(tmp_dir, exist_ok=True)

    html = render_to_string(template_src, context)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)

    script_path = os.path.join(settings.BASE_DIR, "pdfgen", "generate_pdf.js")
    result = subprocess.run(
        ["node", script_path, html_path, pdf_path],
        capture_output=True,
        text=True
    )

    if result.returncode != 0:
        if os.path.exists(html_path):
            os.remove(html_path)
        if os.path.exists(pdf_path):
            os.remove(pdf_path)
        raise Exception(f"Error generando PDF:\n{result.stderr.strip()}")

    if not os.path.exists(pdf_path):
        raise Exception("El archivo PDF no se generó correctamente.")

    if os.path.exists(html_path):
        os.remove(html_path)

    response = FileResponse(open(pdf_path, 'rb'), content_type='application/pdf')
    disposition = 'attachment' if as_attachment else 'inline'
    response['Content-Disposition'] = f'{disposition}; filename="{filename}"'
    return response


def _obtener_rango_fechas_desde_request(request):
    """Devuelve (inicio_dt, fin_dt, inicio_str, fin_str) o None si falta/está mal."""
    fecha_inicio = (request.GET.get('fecha_inicio') or '').strip()
    fecha_fin = (request.GET.get('fecha_fin') or '').strip()

    if not fecha_inicio or not fecha_fin:
        return None

    try:
        inicio_date = datetime.strptime(fecha_inicio, '%Y-%m-%d').date()
        fin_date = datetime.strptime(fecha_fin, '%Y-%m-%d').date()
    except ValueError:
        return None

    if fin_date < inicio_date:
        return None

    tz = timezone.get_current_timezone()
    inicio_dt = timezone.make_aware(datetime.combine(inicio_date, time.min), tz)
    fin_dt = timezone.make_aware(datetime.combine(fin_date, time.max), tz)
    return inicio_dt, fin_dt, fecha_inicio, fecha_fin


def _obtener_ajustes_en_rango(inicio_dt, fin_dt):
    return (
        InventoryAdjustment.objects
        .select_related('item', 'user')
        .prefetch_related('history')
        .filter(created_at__range=(inicio_dt, fin_dt))
        .order_by('created_at', 'id')
    )


@login_required
def exportar_ajustes_inventario_pdf(request):
    rango = _obtener_rango_fechas_desde_request(request)
    if not rango:
        messages.error(request, "Debe seleccionar un rango de fechas válido (inicio y fin).")
        return redirect('inventory_adjustment_create')

    inicio_dt, fin_dt, fecha_inicio, fecha_fin = rango
    ajustes = list(_obtener_ajustes_en_rango(inicio_dt, fin_dt))

    total_entradas = sum(a.quantity for a in ajustes if a.adjustment_type == 'IN')
    total_salidas = sum(a.quantity for a in ajustes if a.adjustment_type == 'OUT')

    context = {
        'logo_url': 'https://raw.githubusercontent.com/Kevin25DC/celupro-assets/refs/heads/main/logo%20de%20prueba%20dos.png',
        'company_name': 'CELUPRO CO',
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
        'ajustes': ajustes,
        'total_registros': len(ajustes),
        'total_entradas': total_entradas,
        'total_salidas': total_salidas,
        'neto': total_entradas - total_salidas,
        'generado_por': (
            request.user.get_full_name().strip()
            if hasattr(request.user, 'get_full_name') and request.user.get_full_name().strip()
            else request.user.username
        ),
        'generado_en': timezone.localtime(timezone.now()),
    }

    filename = f"Reporte_Ajustes_Inventario_{fecha_inicio}_a_{fecha_fin}.pdf"
    return _render_pdf_inline_with_puppeteer(
        'reports/reportsInventory/reporte_ajustes_inventario_rango.html',
        context,
        filename=filename,
        as_attachment=True,
    )


@login_required
def exportar_ajustes_inventario_excel(request):
    rango = _obtener_rango_fechas_desde_request(request)
    if not rango:
        messages.error(request, "Debe seleccionar un rango de fechas válido (inicio y fin).")
        return redirect('inventory_adjustment_create')

    inicio_dt, fin_dt, fecha_inicio, fecha_fin = rango
    ajustes = list(_obtener_ajustes_en_rango(inicio_dt, fin_dt))

    wb = Workbook()
    ws = wb.active
    ws.title = "Ajustes"

    titulo = f"Reporte de ajustes de inventario ({fecha_inicio} a {fecha_fin})"
    ws.merge_cells('A1:I1')
    c = ws['A1']
    c.value = titulo
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal='center')

    headers = [
        'Fecha',
        'Código',
        'Producto',
        'Tipo',
        'Cantidad',
        'Motivo',
        'Usuario',
        'Stock antes',
        'Stock después',
    ]
    ws.append([])
    ws.append(headers)
    header_row = ws[3]
    for cell in header_row:
        cell.font = Font(bold=True)

    total_entradas = 0
    total_salidas = 0

    for adj in ajustes:
        h = None
        try:
            h = adj.history.all()[0]
        except Exception:
            h = None

        tipo = 'Entrada' if adj.adjustment_type == 'IN' else 'Salida'
        cantidad = adj.quantity
        if adj.adjustment_type == 'IN':
            total_entradas += cantidad
        else:
            total_salidas += cantidad

        usuario = (
            adj.user.get_full_name().strip()
            if hasattr(adj.user, 'get_full_name') and adj.user.get_full_name().strip()
            else adj.user.username
        )

        codigo = f"AJI-{timezone.localtime(adj.created_at).strftime('%Y')}-{adj.id:04d}"
        ws.append([
            timezone.localtime(adj.created_at).strftime('%d/%m/%Y %H:%M'),
            codigo,
            getattr(adj.item, 'Name', ''),
            tipo,
            cantidad,
            adj.reason or '',
            usuario,
            h.stock_before if h else '',
            h.stock_after if h else '',
        ])

    ws.append([])
    ws.append(['Totales', '', '', '', '', '', '', '', ''])
    ws.append(['Entradas (unidades)', total_entradas, '', '', '', '', '', '', ''])
    ws.append(['Salidas (unidades)', total_salidas, '', '', '', '', '', '', ''])
    ws.append(['Neto (entradas - salidas)', total_entradas - total_salidas, '', '', '', '', '', '', ''])

    # Ajuste simple de ancho de columnas
    widths = [18, 18, 30, 12, 10, 30, 20, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Reporte_Ajustes_Inventario_{fecha_inicio}_a_{fecha_fin}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def inventory_adjustment_receipt(request, adjustment_id: int):
    adjustment = get_object_or_404(
        InventoryAdjustment.objects.select_related('item', 'user'),
        id=adjustment_id
    )
    history = InventoryHistory.objects.filter(adjustment=adjustment).first()

    adjustment_code = f"AJI-{adjustment.created_at.strftime('%Y')}-{adjustment.id:04d}"

    responsable = (
        adjustment.user.get_full_name().strip()
        if hasattr(adjustment.user, 'get_full_name') and adjustment.user.get_full_name().strip()
        else adjustment.user.username
    )

    context = {
        'logo_url': 'https://raw.githubusercontent.com/Kevin25DC/celupro-assets/refs/heads/main/logo%20de%20prueba%20dos.png',
        'company_name': 'CELUPRO CO',
        'adjustment': adjustment,
        'history': history,
        'adjustment_code': adjustment_code,
        'responsable': responsable,
    }

    filename = f"Recibo_Ajuste_Inventario_{adjustment_code}.pdf"
    return _render_pdf_inline_with_puppeteer(
        'reports/reportsInventory/ajusteInventario.html',
        context,
        filename=filename
    )
